"""Offline tests: no Ollama, microphone or network needed (embeddings are faked)."""
from __future__ import annotations

import hashlib
import os
import sys
import tempfile
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import numpy as np
import pytest

_TMP = tempfile.mkdtemp(prefix="foresight-test-")
os.environ["FORESIGHT_DATA_DIR"] = _TMP
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from foresight import (assistant, briefings, calendars, db, indexer, llm, media, media_index, meetings, parsers,  # noqa: E402
                       repass, retrieval, runtime)

DIM = 64
media_index._started = True  # no background audio indexing (it would call Ollama); tests call index_media() directly


def fake_embed(texts, *, query=False, title="none", batch=16):
    """Bag-of-words hashing embedding: shared words → high cosine similarity."""
    out = np.zeros((len(texts), DIM), dtype=np.float32)
    for i, text in enumerate(texts):
        for word in text.lower().split():
            word = word.strip(".,?!:;()[]\"'")
            if len(word) > 2:
                out[i, int(hashlib.md5(word.encode()).hexdigest(), 16) % DIM] += 1.0
    norms = np.linalg.norm(out, axis=1, keepdims=True)
    norms[norms == 0] = 1
    return out / norms


@pytest.fixture(autouse=True)
def _offline(monkeypatch):
    monkeypatch.setattr(llm, "embed", fake_embed)
    # Test audio is synthetic tones, which a real speech detector rightly rejects: count all of it as speech.
    monkeypatch.setattr(repass, "speech_seconds", lambda pcm: len(pcm) / 16000)
    # Gemma re-ranking is one extra model call per answer; tests that count calls opt in to it explicitly.
    monkeypatch.setitem(runtime._load(), "rerank", False)

    def no_gemma(*args, **kwargs):  # the real model would make tests slow and flaky: fake it where a test needs it
        raise llm.LLMError("offline test: fake llm.chat / llm.chat_stream in this test")

    monkeypatch.setattr(llm, "chat", no_gemma)
    monkeypatch.setattr(llm, "chat_stream", no_gemma)
    db.init()


def test_chunking_respects_size_and_overlap():
    text = "\n\n".join(f"Paragraph {i}. " + "Sentence about the budget. " * 12 for i in range(10))
    chunks = indexer.chunk_text(text, size=500, overlap=100)
    assert len(chunks) > 5
    assert all(len(c) <= 600 for c in chunks)
    assert indexer.chunk_text("   ") == []


def test_parsers_docx_xlsx_html(tmp_path):
    import docx
    from openpyxl import Workbook

    d = docx.Document()
    d.add_heading("Launch plan", level=1)
    d.add_paragraph("The launch budget is $40,000.")
    d.save(tmp_path / "plan.docx")
    assert "# Launch plan" in parsers.extract(tmp_path / "plan.docx")
    assert "$40,000" in parsers.extract(tmp_path / "plan.docx")

    wb = Workbook()
    wb.active.append(["Item", "Cost"])
    wb.active.append(["Ads", 1200])
    wb.save(tmp_path / "costs.xlsx")
    assert "Ads | 1200" in parsers.extract(tmp_path / "costs.xlsx")

    (tmp_path / "page.html").write_text("<html><title>T</title><body><script>x()</script><h2>Hi</h2><p>Body text</p></body></html>")
    out = parsers.extract(tmp_path / "page.html")
    assert "## Hi" in out and "Body text" in out and "x()" not in out

    (tmp_path / "bad.pdf").write_bytes(b"not a pdf")
    with pytest.raises(parsers.ParseError):
        parsers.extract(tmp_path / "bad.pdf")


def test_question_detection():
    found = assistant.candidate_questions("Okay let's start. What was the budget for the Q3 launch? I think it moved.")
    assert found == ["What was the budget for the Q3 launch?"]
    assert assistant.candidate_questions("Yes.") == []
    assert assistant.candidate_questions("how many vendors signed the contract") == ["how many vendors signed the contract"]


def test_action_item_parsing():
    md = "## Summary\nx\n## Action items\n- [ ] Priya — send vendor list (Friday)\n- [ ] Unassigned — book room\n## Other\n- no"
    assert meetings._parse_action_items(md) == ["Priya — send vendor list (Friday)", "Unassigned — book room"]
    assert meetings._parse_action_items("## Action items\n- None recorded.") == []


def test_ingest_and_hybrid_search(tmp_path):
    (tmp_path / "budget.md").write_text("# Budget\nThe marketing budget for the Atlas launch is 40000 dollars.", encoding="utf-8")
    (tmp_path / "team.txt").write_text("Priya owns the vendor shortlist. Marco runs the design review.", encoding="utf-8")
    folder = indexer.add_folder(str(tmp_path))
    assert folder["queued"] == 2
    while not indexer._queue.empty():
        sid = indexer._queue.get()
        indexer._pending.discard(sid)
        indexer._process(sid)
    ready = db.query("SELECT title, status, n_chunks FROM sources WHERE folder_id = ?", (folder["id"],))
    assert {r["status"] for r in ready} == {"ready"}

    hits = retrieval.search("what is the marketing budget for Atlas", k=3)
    assert hits and hits[0].title == "budget.md"
    assert hits[0].similarity > 0.3
    hits = retrieval.search("who owns the vendor shortlist", k=3)
    assert hits[0].title == "team.txt"

    # Labels restrict retrieval to tagged sources.
    sid = db.query_one("SELECT id FROM sources WHERE title = 'team.txt'")["id"]
    with db.tx() as conn:
        conn.execute("INSERT INTO labels(name) VALUES ('People')")
        conn.execute("INSERT INTO source_labels VALUES (?, (SELECT id FROM labels WHERE name='People'))", (sid,))
    assert {h.title for h in retrieval.search("budget vendor", k=5, label="People")} == {"team.txt"}

    # An unchanged file is not re-embedded.
    before = db.query_one("SELECT indexed_at FROM sources WHERE id = ?", (sid,))["indexed_at"]
    indexer._process(sid)
    assert db.query_one("SELECT indexed_at FROM sources WHERE id = ?", (sid,))["indexed_at"] == before


def test_local_only_guard():
    from fastapi.testclient import TestClient

    from foresight.server import app

    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        assert client.get("/api/labels").status_code == 200
        assert client.put("/api/settings", json={"live_assist": False}).status_code == 403
        assert client.put("/api/settings", json={"live_assist": False}, headers={"X-Foresight": "1"}).json()["live_assist"] is False
        for path in ("/", "/static/app.js"):  # the UI is revalidated, so an update is never hidden by the cache
            assert client.get(path).headers["cache-control"] == "no-cache"
    with TestClient(app, base_url="http://evil.example:8765") as client:
        assert client.get("/api/labels").status_code == 403


ICS_FIXTURE = r"""BEGIN:VCALENDAR
VERSION:2.0
X-WR-CALNAME:Team
BEGIN:VEVENT
UID:weekly-1
SUMMARY:Atlas weekly sync\, budget
DTSTART;TZID=America/New_York:20260921T100000
DTEND;TZID=America/New_York:20260921T103000
RRULE:FREQ=WEEKLY;BYDAY=MO,WE,FR;UNTIL=20261231T000000Z
EXDATE;TZID=America/New_York:20261021T100000
ATTENDEE;CN="Kim, Dana";ROLE=REQ-PARTICIPANT:mailto:dana@example.com
ATTENDEE;CUTYPE=ROOM;CN=Room 4:mailto:room4@example.com
ATTENDEE:mailto:priya@example.com
DESCRIPTION:First line\nsecond line that is folded
  onto two lines
BEGIN:VALARM
TRIGGER:-PT10M
DESCRIPTION:not the event description
END:VALARM
END:VEVENT
BEGIN:VEVENT
UID:weekly-1
RECURRENCE-ID;TZID=America/New_York:20261026T100000
SUMMARY:Atlas weekly sync (moved)
DTSTART;TZID=America/New_York:20261026T140000
DURATION:PT45M
END:VEVENT
BEGIN:VEVENT
UID:daily-1
SUMMARY:Standup
DTSTART:20260101T150000Z
DURATION:PT15M
RRULE:FREQ=DAILY;COUNT=400;BYDAY=MO,TU,WE,TH,FR
END:VEVENT
BEGIN:VEVENT
UID:once-1
SUMMARY:Vendor call
DTSTART:20261023T090000Z
DTEND:20261023T100000Z
END:VEVENT
BEGIN:VEVENT
UID:cancelled-1
SUMMARY:Cancelled thing
STATUS:CANCELLED
DTSTART:20261022T090000Z
DTEND:20261022T100000Z
END:VEVENT
BEGIN:VEVENT
UID:later-1
SUMMARY:Outside the window
DTSTART:20261110T090000Z
DTEND:20261110T100000Z
END:VEVENT
END:VCALENDAR
"""


def _utc(*args) -> float:
    return datetime(*args, tzinfo=timezone.utc).timestamp()


def test_ics_parsing_and_recurrence_expansion():
    week = calendars.parse_ics(ICS_FIXTURE, start=_utc(2026, 10, 20), end=_utc(2026, 10, 27))
    got = [(datetime.fromtimestamp(e.starts_at, timezone.utc).strftime("%a %H:%M"), int(e.ends_at - e.starts_at) // 60, e.title)
           for e in week]
    assert got == [
        ("Tue 15:00", 15, "Standup"),
        ("Wed 15:00", 15, "Standup"),            # Wed 21 Oct weekly instance is an EXDATE
        ("Thu 15:00", 15, "Standup"),
        ("Fri 09:00", 60, "Vendor call"),        # the cancelled event and the one in November are left out
        ("Fri 14:00", 30, "Atlas weekly sync, budget"),
        ("Fri 15:00", 15, "Standup"),
        ("Mon 15:00", 15, "Standup"),            # weekdays only: no weekend standups
        ("Mon 18:00", 45, "Atlas weekly sync (moved)"),  # RECURRENCE-ID moved 10:00 → 14:00 New York time
    ]
    assert len({e.uid for e in week}) == len(week)
    weekly = next(e for e in week if e.title.startswith("Atlas weekly sync,"))
    assert weekly.attendees == ["Kim, Dana", "priya@example.com"]  # rooms are not people
    assert weekly.description == "First line\nsecond line that is folded onto two lines"

    # The weekly rule repeats on New York wall-clock time: after DST ends (1 Nov) 10:00 is 15:00 UTC.
    november = calendars.parse_ics(ICS_FIXTURE, start=_utc(2026, 11, 2), end=_utc(2026, 11, 9))
    assert [datetime.fromtimestamp(e.starts_at, timezone.utc).strftime("%a %d %H:%M")
            for e in november if e.uid.startswith("weekly-1")] == ["Mon 02 15:00", "Wed 04 15:00", "Fri 06 15:00"]

    with pytest.raises(calendars.CalendarError):
        calendars.parse_ics("hello", start=0, end=1)


def _write_event_ics(path: Path, title: str, start: datetime, *, attendee: str = "Lena Ortiz") -> None:
    path.write_text(
        "BEGIN:VCALENDAR\r\nVERSION:2.0\r\nBEGIN:VEVENT\r\nUID:harbor-1\r\n"
        f"SUMMARY:{title}\r\nDTSTART:{start.strftime('%Y%m%dT%H%M%SZ')}\r\nDURATION:PT30M\r\n"
        "DESCRIPTION:Review the kiosk rollout budget. Join https://meet.example.com/abc\r\n"
        f"ATTENDEE;CN={attendee}:mailto:lena@example.com\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n", encoding="utf-8")


def test_calendar_sync_and_briefing_cache(tmp_path, monkeypatch):
    doc = tmp_path / "harbor_budget.md"
    doc.write_text("The Harbor kiosk rollout budget is 75000 dollars. Lena Ortiz owns the rollout.", encoding="utf-8")
    # Index synchronously; going through the queue would race the worker thread an earlier TestClient started.
    indexer._process(indexer._upsert_source(str(doc), doc.name, "upload", mime="md"))

    soon = datetime.now(timezone.utc).replace(microsecond=0) + timedelta(hours=2)
    ics = tmp_path / "work.ics"
    _write_event_ics(ics, "Harbor kiosk budget review", soon)
    cal = calendars.add_calendar(str(ics))
    assert cal["n_events"] == 1
    past = db.execute("INSERT INTO meetings(title, started_at, ended_at, status, summary) VALUES (?,?,?,?,?)",
                      ("Harbor kiosk budget sync", time.time() - 7 * 86400, time.time() - 7 * 86400 + 1800, "done",
                       "Lena asked the vendor for a revised quote."))

    calls = []

    def fake_chat(messages, **kwargs):
        calls.append(messages)
        return "Here you go:\n- The budget is 75000 dollars [1]\n- Lena Ortiz owns it [1]\n- A revised quote is pending [2]\n- extra"

    monkeypatch.setattr(llm, "chat", fake_chat)
    try:
        [ev] = briefings.upcoming(enqueue_missing=False)
        assert ev["status"] == "pending" and ev["attendees"] == ["Lena Ortiz"]
        assert "meet.example.com" not in briefings.query_text(briefings._event(ev["id"]))

        b = briefings.generate(ev["id"])
        assert len(calls) == 1 and "Lena Ortiz" in calls[0][1]["content"]
        assert b["summary"].splitlines() == ["- The budget is 75000 dollars [1]", "- Lena Ortiz owns it [1]",
                                             "- A revised quote is pending [2]"]
        assert b["citations"][0]["title"] == "harbor_budget.md"
        assert [c["title"] for c in b["citations"] if c.get("meeting_id") == past] == [
            f"Harbor kiosk budget sync ({time.strftime('%a %Y-%m-%d', time.localtime(time.time() - 7 * 86400))})"]

        # Unchanged event (even after a re-sync): served from the cache, no second Gemma call.
        calendars.sync(cal["id"])
        [ev2] = briefings.upcoming(enqueue_missing=False)
        assert ev2["id"] == ev["id"] and ev2["status"] == "ready"
        briefings.generate(ev["id"])
        assert len(calls) == 1

        # A changed event keeps its id but needs a new briefing.
        _write_event_ics(ics, "Harbor kiosk budget review", soon, attendee="Lena Ortiz-Park")
        calendars.sync(cal["id"])
        [ev3] = briefings.upcoming(enqueue_missing=False)
        assert ev3["id"] == ev["id"] and ev3["status"] == "pending"
        briefings.generate(ev["id"])
        assert len(calls) == 2
    finally:
        calendars.remove_calendar(cal["id"])
        db.execute("DELETE FROM meetings WHERE id = ?", (past,))
    assert db.query_one("SELECT COUNT(*) AS n FROM events")["n"] == 0
    assert db.query_one("SELECT COUNT(*) AS n FROM briefings")["n"] == 0


def test_calendar_api(tmp_path):
    from fastapi.testclient import TestClient

    from foresight.server import app

    ics = tmp_path / "later.ics"
    _write_event_ics(ics, "Quarterly planning", datetime.now(timezone.utc) + timedelta(days=3))
    (tmp_path / "notes.ics").write_text("just some notes", encoding="utf-8")
    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        assert client.post("/api/calendars", json={"uri": str(ics)}).status_code == 403
        added = client.post("/api/calendars", json={"uri": str(ics)}, headers={"X-Foresight": "1"})
        assert added.status_code == 200 and added.json()["n_events"] == 1
        for bad in (tmp_path / "missing.ics", tmp_path / "notes.ics", "ftp://example.com/cal.ics"):
            assert client.post("/api/calendars", json={"uri": str(bad)}, headers={"X-Foresight": "1"}).status_code == 400
        assert client.get("/api/briefings").json() == {"calendars": 1, "events": []}  # 3 days away: not in the next 24 h
        assert client.delete(f"/api/calendars/{added.json()['id']}").status_code == 403
        assert client.delete(f"/api/calendars/{added.json()['id']}", headers={"X-Foresight": "1"}).status_code == 200
        assert client.get("/api/calendars").json() == []


# ---------------------------------------------------------------- stored audio (multimodal memory, P0–P1)

def _tone(seconds: float, freq: float = 220.0) -> np.ndarray:
    t = np.arange(int(seconds * media.RATE)) / media.RATE
    return (0.2 * np.sin(2 * np.pi * freq * t)).astype(np.float32)


@pytest.fixture
def keep_audio():
    """Turn "Keep meeting audio" on for one test, and restore the defaults afterwards."""
    runtime.update({"keep_meeting_audio": True, "audio_retention_days": 30})
    yield
    runtime.update({"keep_meeting_audio": False, "audio_retention_days": 30})


def _new_meeting(title: str = "Audio test") -> int:
    return db.execute("INSERT INTO meetings(title, started_at, ended_at, status) VALUES (?,?,?,?)",
                      (title, time.time(), time.time() + 60, "done"))


def test_wav_bytes_and_media_capability_check(monkeypatch):
    import io
    import wave

    with wave.open(io.BytesIO(llm.wav_bytes(_tone(0.5)))) as w:
        assert (w.getnchannels(), w.getframerate(), w.getsampwidth(), w.getnframes()) == (1, 16000, 2, 8000)

    class Show:
        def post(self, path, json=None, timeout=None):
            caps = ["embedding", "vision", "audio"] if "embedding" in json["model"] else ["completion", "vision", "audio"]
            return type("R", (), {"json": staticmethod(lambda: {"capabilities": caps})})()

    monkeypatch.setattr(llm, "_client", Show())
    monkeypatch.setattr(llm, "embed_audio", lambda wavs: np.eye(2, dtype=np.float32))
    monkeypatch.setattr(llm, "embed_images", lambda pictures: np.eye(2, dtype=np.float32))
    caps = llm.check_media()
    assert caps["audio_embed"] and caps["image_embed"] and caps["gen_audio"] and caps["gen_vision"] and caps["error"] is None
    # The silent failure: Ollama ignores the audio and returns the same vector for every clip.
    monkeypatch.setattr(llm, "embed_audio", lambda wavs: np.ones((2, 4), dtype=np.float32) / 2)
    monkeypatch.setattr(llm, "embed_images", lambda pictures: np.ones((2, 4), dtype=np.float32) / 2)
    caps = llm.check_media()
    assert not caps["audio_embed"] and not caps["image_embed"] and "identical" in caps["error"]
    assert llm.media_status() == caps


def test_audio_store_retention_and_cleanup():
    mid = _new_meeting()
    store = media.AudioStore(mid, ["You", "Others"])
    for block in np.split(_tone(3.0), 12):            # 3 s from the microphone, in 250 ms capture blocks
        store.feed("You", block)
    store.feed("Others", _tone(2.0, 330))
    store.writers["Others"]._q.put_nowait(media.RATE)  # 1 s the encoder could not keep up with: written as silence
    store.feed("Unknown", _tone(1.0))                  # channels that were not opened are ignored
    store.close()

    tracks = {t["channel"]: t for t in media.for_meeting(mid)}
    assert set(tracks) == {"You", "Others"} and all(t["status"] == "ready" and t["url"] for t in tracks.values())
    assert tracks["You"]["duration_s"] == pytest.approx(3.0) and tracks["Others"]["duration_s"] == pytest.approx(3.0)
    paths = {r["channel"]: Path(r["path"]) for r in db.query("SELECT channel, path FROM media WHERE meeting_id = ?", (mid,))}
    assert media.probe_duration(paths["Others"]) == pytest.approx(3.0, abs=0.1)  # the gap keeps the timeline aligned

    stray = media.MEDIA_DIR / "stray.ogg"
    stray.write_bytes(b"x")
    db.execute("UPDATE media SET created_at = ? WHERE meeting_id = ? AND channel = 'You'", (time.time() - 10 * 86400, mid))
    runtime.update({"audio_retention_days": 7})
    try:
        media.apply_retention()  # the 10-day-old file is past a 7-day retention; the orphan goes too
    finally:
        runtime.update({"audio_retention_days": 30})
    tracks = {t["channel"]: t for t in media.for_meeting(mid)}
    assert tracks["You"]["status"] == "expired" and tracks["You"]["url"] is None and not paths["You"].exists()
    assert tracks["Others"]["status"] == "ready" and paths["Others"].exists() and not stray.exists()

    meetings.delete(mid)
    assert not paths["Others"].exists() and not paths["Others"].parent.exists()
    assert db.query("SELECT id FROM media WHERE meeting_id = ?", (mid,)) == []


def test_imported_recording_is_kept_only_when_audio_is_kept(keep_audio):
    from foresight.config import settings

    def recording(name: str) -> Path:
        path = settings.recordings_dir / name
        path.write_bytes(llm.wav_bytes(_tone(2.0)))
        return path

    kept, mid = recording("kept.wav"), _new_meeting("Imported")
    media_id = media.adopt_import(mid, kept)
    row = db.query_one("SELECT * FROM media WHERE id = ?", (media_id,))
    assert row["status"] == "ready" and row["mime"] == "audio/wav" and row["duration_s"] == pytest.approx(2.0, abs=0.05)
    assert row["expires_at"] == pytest.approx(row["created_at"] + 30 * 86400)

    runtime.update({"keep_meeting_audio": False})
    dropped = recording("dropped.wav")
    assert media.adopt_import(_new_meeting("Imported 2"), dropped) is None and not dropped.exists()

    meetings.delete(mid)
    assert not kept.exists() and settings.recordings_dir.is_dir()  # the shared recordings folder itself stays


def test_media_api_playback_and_delete(keep_audio):
    from fastapi.testclient import TestClient

    from foresight.config import settings
    from foresight.server import app

    path = settings.recordings_dir / "api.wav"
    path.write_bytes(llm.wav_bytes(_tone(1.0)))
    mid = _new_meeting("Playback")
    media.adopt_import(mid, path)
    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        [track] = client.get(f"/api/meetings/{mid}").json()["media"]
        part = client.get(track["url"], headers={"Range": "bytes=0-99"})
        assert part.status_code == 206 and len(part.content) == 100 and part.headers["content-type"] == "audio/wav"
        assert client.delete("/api/media").status_code == 403
        assert client.delete("/api/media", headers={"X-Foresight": "1"}).json() == {"deleted": 1}
        assert client.get(track["url"]).status_code == 404
        assert client.get(f"/api/meetings/{mid}").json()["media"] == []
    assert not path.exists()
    meetings.delete(mid)


def test_eval_scoring_helpers():
    import importlib.util

    spec = importlib.util.spec_from_file_location("eval_media", Path(__file__).resolve().parents[1] / "tools" / "eval_media.py")
    ev = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ev)

    assert ev.wer("The budget is $52,000.", "the budget is $52000") == 0
    assert ev.wer("one two three four", "one three four five") == pytest.approx(0.5)
    clips = [{"id": "a"}, {"id": "b"}, {"id": "c"}]
    questions = [{"q": "budget", "relevant": {"b"}}, {"q": "venue", "relevant": {"c"}}]
    # The second list counts double, so its top results win where the two lists disagree.
    assert ev.fuse([([[0, 1, 2], [2, 1, 0]], 1.0), ([[1, 0, 2], [2, 0, 1]], 2.0)], 2) == [[1, 0, 2], [2, 0, 1]]
    assert ev.metrics([[1, 0, 2], [0, 2, 1]], questions, clips) == {"recall@1": 0.5, "recall@5": 1.0, "mrr": 0.75}
    assert ev.rank_by_keywords([{"q": "what is the venue cost"}], ["budget talk", "venue costs 11800", "hiring"]) == [[1]]


# ---------------------------------------------------------------- audio moments in search (P2)

def _segs(*rows):
    return [{"t_start": a, "t_end": b, "text": t} for a, b, t in rows]


def test_build_windows():
    # ~20 s windows that share one line, so a sentence on a boundary is whole in one of them
    w = media_index.build_windows(_segs((0, 8, "a"), (9, 17, "b"), (18, 26, "c"), (27, 35, "d")))
    assert [(x.t_start, x.t_end, x.text) for x in w] == [(0, 26, "a b c"), (18, 35, "c d")]
    # a pause longer than 5 s starts a new window, without overlap
    w = media_index.build_windows(_segs((0, 6, "a"), (7, 12, "b"), (30, 36, "c")))
    assert [(x.t_start, x.t_end) for x in w] == [(0, 12), (30, 36)]
    # one long line is one window; nothing in, nothing out
    assert [(x.t_start, x.t_end) for x in media_index.build_windows(_segs((0, 40, "x")))] == [(0, 40)]
    assert media_index.build_windows([]) == []


STANDUP = _segs(
    (0.0, 6.0, "After the finance review the Atlas marketing budget was raised to 52000 dollars"),
    (7.0, 12.0, "The extra money goes to a second wave of paid social in October"),
    (13.0, 19.0, "Priya needs final quotes from the shortlisted vendors by Friday"),
    (30.0, 36.0, "The printing vendor needs fifteen business days of lead time"),
)


def test_audio_moments_are_indexed_and_found(keep_audio, monkeypatch):
    from fastapi.testclient import TestClient

    from foresight.config import settings
    from foresight.server import app

    mid = _new_meeting("Imported standup")
    with db.tx() as conn:
        conn.executemany("INSERT INTO segments(meeting_id, t_start, t_end, speaker, text) VALUES (?,?,?,?,?)",
                         [(mid, s["t_start"], s["t_end"], "Recording", s["text"]) for s in STANDUP])
    path = settings.recordings_dir / "standup.wav"
    path.write_bytes(llm.wav_bytes(_tone(40.0)))
    media_id = media.adopt_import(mid, path)

    # Offline stand-in for EmbeddingGemma 2 audio: the vector of the words spoken in each window, in order.
    spoken = iter([w.text for w in media_index.build_windows(STANDUP)])
    monkeypatch.setattr(llm, "embed_audio", lambda wavs: fake_embed([next(spoken) for _ in wavs]))
    try:
        assert media_index.index_media(media_id) == 2
        assert media_index.index_media(media_id) == 0  # resumes from its progress, never stores a window twice
        assert media_index.pending() == [] and media.usage()["moments"] == 2

        # Without an indexed transcript, the moment itself is the result, playable from where it was said.
        hits = retrieval.search("who needs quotes from the vendors by Friday", k=6)
        moment = next(h for h in hits if h.kind == "audio")
        assert moment.media == {"media_id": media_id, "channel": "Recording", "t_start": 0.0, "t_end": 19.0,
                                "url": f"/api/media/{media_id}"}
        assert moment.title == "Imported standup — audio 00:00" and "Priya" in moment.text
        assert moment.as_citation(1)["media"]["url"] == f"/api/media/{media_id}"
        assert not any(h.media for h in retrieval.search("vendors by Friday", k=6, with_media=False))
        assert not any(h.media for h in retrieval.search("vendors by Friday", k=6, exclude_meeting=mid))

        # With the transcript indexed, the moment becomes the play button of the transcript result (one card).
        sid = indexer._upsert_source(f"meeting:{mid}:transcript", "Imported standup — transcript", "transcript",
                                     meeting_id=mid, mime="md")
        indexer._process(sid)
        hits = retrieval.search("Priya needs final quotes from the shortlisted vendors by Friday", k=6)
        merged = [h for h in hits if h.kind == "transcript" and h.source_id == sid]
        assert merged and merged[0].media["t_start"] == 0.0
        assert not any(h.kind == "audio" and h.media["t_start"] == 0.0 for h in hits)

        with TestClient(app, base_url="http://127.0.0.1:8765") as client:
            found = client.get("/api/search", params={"q": "printing vendor lead time", "modality": "audio"}).json()
            assert found and all(c["media"]["url"] == f"/api/media/{media_id}" for c in found)
            assert not any("media" in c for c in client.get("/api/search", params={"q": "printing vendor lead time",
                                                                                     "modality": "text"}).json())
    finally:
        meetings.delete(mid)
    # Deleting the meeting removes its audio moments from the index too.
    assert db.query("SELECT id FROM media_items") == []
    assert not any(h.media for h in retrieval.search("printing vendor lead time", k=6))


# ---------------------------------------------------------------- Gemma listens (P3)

def _standup_meeting(title: str = "Listening standup") -> tuple[int, int, list[int]]:
    """A meeting with the STANDUP lines and a stored 40 s recording. Returns (meeting, media, segment ids)."""
    from foresight.config import settings

    mid = _new_meeting(title)
    with db.tx() as conn:
        seg_ids = [conn.execute("INSERT INTO segments(meeting_id, t_start, t_end, speaker, text) VALUES (?,?,?,?,?)",
                                (mid, s["t_start"], s["t_end"], "Recording", s["text"])).lastrowid for s in STANDUP]
    path = settings.recordings_dir / f"listen-{mid}.wav"
    path.write_bytes(llm.wav_bytes(_tone(40.0)))
    return mid, media.adopt_import(mid, path), seg_ids


def test_decode_slice_matches_full_decode(tmp_path):
    rng = np.random.default_rng(3)
    path = tmp_path / "noise.wav"
    path.write_bytes(llm.wav_bytes((0.3 * rng.standard_normal(16000 * 20)).astype(np.float32)))
    full = media.decode_pcm16(path)
    for t0, t1 in ((0.0, 2.0), (7.25, 9.5), (18.0, 25.0)):
        part = media.decode_slice(path, t0, t1)
        assert np.array_equal(part, full[int(t0 * 16000):int(min(t1, 20.0) * 16000)])


def test_chat_hears_the_picked_moment_and_cited_audio(keep_audio, monkeypatch):
    from foresight import listen, rag

    mid, media_id, _ = _standup_meeting()
    spoken = iter([w.text for w in media_index.build_windows(STANDUP)])
    monkeypatch.setattr(llm, "embed_audio", lambda wavs: fake_embed([next(spoken) for _ in wavs]))
    media_index.index_media(media_id)
    sent = []

    def fake_stream(messages, **kwargs):
        sent.append(messages)
        yield "Priya needs the quotes [1]."

    monkeypatch.setattr(llm, "chat_stream", fake_stream)
    hearings = []  # Gemma transcribing each clip on its own before answering
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: hearings.append(messages[-1]["images"])
                        or "Priya needs final quotes from two vendors by Friday.")
    try:
        # "Ask about this moment": the clip and its transcript go to Gemma, cited as one more numbered block.
        events = list(rag.answer_stream("What does Priya need?", focus={"media_id": media_id, "t_start": 12.0, "t_end": 20.0}))
        start = events[0]
        picked = start["citations"][-1]
        assert start["listened"] == [picked["n"]] and picked["kind"] == "audio"
        assert picked["media"] == {"media_id": media_id, "channel": "Recording", "t_start": 12.0, "t_end": 20.0,
                                   "url": f"/api/media/{media_id}"}
        assert "Priya needs final quotes" in picked["snippet"]
        user = sent[-1][-1]
        assert len(hearings) == 1 and len(hearings[0]) == 1 and base64_wav(hearings[0][0]).startswith(b"RIFF")
        assert "images" not in user  # heard once; the answer works from the words
        assert f"[{picked['n']}]" in user["content"] and "the moment the user is asking about" in user["content"]
        assert "Answer from what was heard" in user["content"]
        assert {"type": "heard", "texts": {str(picked["n"]): "Priya needs final quotes from two vendors by Friday."}} in events
        assert "Heard in the recording, word for word: Priya needs final quotes" in user["content"]

        # "Listen to cited audio": cited moments are heard too, at most two clips, none without stored audio.
        before = len(hearings)
        events = list(rag.answer_stream("when are the vendor quotes due", listen_audio=True))
        heard = events[0]["listened"]
        assert 1 <= len(heard) <= rag.MAX_CLIPS and len(hearings) - before == len(heard)
        assert all(c.get("media") for c in events[0]["citations"] if c["n"] in heard)
        # Without either, Gemma gets text only.
        list(rag.answer_stream("when are the vendor quotes due"))
        assert "images" not in sent[-1][-1]
        assert listen.clip_wav(media_id, 50.0, 60.0) is None  # past the end of the recording
    finally:
        meetings.delete(mid)


def base64_wav(field: str) -> bytes:
    import base64
    return base64.b64decode(field)


def test_retranscribe_and_correct_a_line(keep_audio, monkeypatch):
    from fastapi.testclient import TestClient

    from foresight import listen
    from foresight.server import app

    mid, media_id, seg_ids = _standup_meeting("Correction standup")
    spoken = iter([w.text for w in media_index.build_windows(STANDUP)])
    monkeypatch.setattr(llm, "embed_audio", lambda wavs: fake_embed([next(spoken) for _ in wavs]))
    media_index.index_media(media_id)
    indexer._upsert_source(f"meeting:{mid}:transcript", "Correction standup — transcript", "transcript",
                           meeting_id=mid, mime="md")
    heard = []
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: heard.append(messages) or "Priya needs final quotes from Dana by Friday")
    reindexed = []
    monkeypatch.setattr(indexer, "add_meeting_text", lambda *a: reindexed.append(a))
    try:
        suggestion = listen.retranscribe(seg_ids[2])
        assert suggestion == {"id": seg_ids[2], "original": STANDUP[2]["text"], "text": "Priya needs final quotes from Dana by Friday"}
        assert len(heard[0][0]["images"]) == 1 and "Transcribe" in heard[0][0]["content"]
        assert db.query_one("SELECT text FROM segments WHERE id = ?", (seg_ids[2],))["text"] == STANDUP[2]["text"]  # not saved yet

        with TestClient(app, base_url="http://127.0.0.1:8765") as client:
            assert client.post(f"/api/segments/{seg_ids[2]}/retranscribe").status_code == 403
            assert client.post(f"/api/segments/{seg_ids[2]}/retranscribe", headers={"X-Foresight": "1"}).json()["text"] == suggestion["text"]
            assert client.post("/api/segments/999999/retranscribe", headers={"X-Foresight": "1"}).status_code == 404
            saved = client.put(f"/api/segments/{seg_ids[2]}", json={"text": "  Priya needs final quotes from Dana by Friday "},
                               headers={"X-Foresight": "1"})
            assert saved.json() == {"id": seg_ids[2], "text": "Priya needs final quotes from Dana by Friday"}
            assert client.put(f"/api/segments/{seg_ids[2]}", json={"text": "  "}, headers={"X-Foresight": "1"}).status_code == 400

            # A line on a channel whose audio is not stored cannot be re-transcribed.
            other = db.execute("INSERT INTO segments(meeting_id, t_start, t_end, speaker, text) VALUES (?,?,?,?,?)",
                               (mid, 1.0, 2.0, "You", "hello"))
            assert client.post(f"/api/segments/{other}/retranscribe", headers={"X-Foresight": "1"}).status_code == 409

        assert db.query_one("SELECT text FROM segments WHERE id = ?", (seg_ids[2],))["text"].endswith("from Dana by Friday")
        window = db.query_one("SELECT text FROM media_items WHERE media_id = ? AND t_start = 0", (media_id,))
        assert "quotes from Dana by Friday" in window["text"] and "Atlas marketing budget" in window["text"]
        assert reindexed == [(mid, "transcript", "Correction standup — transcript")]
    finally:
        meetings.delete(mid)


def test_eval_handles_long_hinglish_recordings(tmp_path):
    import importlib.util
    import json

    spec = importlib.util.spec_from_file_location("eval_media", Path(__file__).resolve().parents[1] / "tools" / "eval_media.py")
    ev = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ev)

    (tmp_path / "standup.wav").write_bytes(llm.wav_bytes(_tone(45.0)))
    (tmp_path / "standup.txt").write_text("budget kitna approve hua aur vendor list Priya ke paas hai", encoding="utf-8")
    (tmp_path / "dataset.json").write_text(json.dumps({
        "recordings": [{"file": "standup.wav", "reference_file": "standup.txt"}],
        "questions": [{"q": "Budget kitna approve hua?", "at": {"file": "standup.wav", "t": 19.0}},
                      {"q": "Who has the vendor list?", "at": {"file": "standup.wav", "t": "0:25", "until": "0:41"}},
                      {"q": "Asked after the end", "at": {"file": "standup.wav", "t": 100}}]}), encoding="utf-8")
    clips, questions, references = ev.load(tmp_path)
    assert [c["id"] for c in clips] == ["standup.wav@0", "standup.wav@20", "standup.wav@40"]
    # ±2 s around a timestamp: an answer spoken across a window boundary counts in either window
    assert [q["relevant"] for q in questions] == [{"standup.wav@0", "standup.wav@20"}, {"standup.wav@20", "standup.wav@40"}]
    assert references == {"standup.wav": "budget kitna approve hua aur vendor list Priya ke paas hai"}

    latin = ["budget kitna approve hua", "aur vendor list Priya", "ke paas hai"]
    assert ev.engine_wer(clips, latin, references) == {"wer": 0.0}
    devanagari = ["बजट कितना अप्रूव हुआ", "और वेंडर लिस्ट प्रिया", "के पास है"]  # same words, other script
    assert ev.devanagari_share(" ".join(devanagari)) == 1.0 and ev.devanagari_share(" ".join(latin)) == 0.0
    assert ev.engine_wer(clips, devanagari, references) == {"wer": None, "note": "1 of 1 transcripts are in another script than the reference"}
    assert ev.wer("बजट कितना है।", "बजट कितना है") == 0  # the danda is punctuation


# ---------------------------------------------------------------- Hindi / Hinglish meetings

def test_migration_adds_meeting_language_to_old_databases(tmp_path):
    import sqlite3

    conn = sqlite3.connect(tmp_path / "old.db")
    conn.row_factory = sqlite3.Row
    conn.execute("CREATE TABLE meetings (id INTEGER PRIMARY KEY, title TEXT NOT NULL, started_at REAL NOT NULL)")
    conn.execute("INSERT INTO meetings(title, started_at) VALUES ('old meeting', 0)")
    db._migrate(conn)
    db._migrate(conn)  # running it again is harmless
    assert conn.execute("SELECT language FROM meetings").fetchone()["language"] == "en"


def test_engine_follows_the_meeting_language(monkeypatch):
    from foresight import asr

    class FakeWhisper:
        def __init__(self, language=None):
            self.language = language

    english = object()
    monkeypatch.setattr(asr, "WhisperEngine", FakeWhisper)
    monkeypatch.setattr(asr.Transcriber, "_instance", english)
    monkeypatch.setattr(asr.Transcriber, "_hindi", None)
    hindi = asr.Transcriber.for_language("hi")
    assert isinstance(hindi, FakeWhisper) and hindi.language == "hi" and asr.Transcriber.for_language("hi") is hindi
    assert asr.Transcriber.for_language("en") is english and asr.Transcriber.for_language() is english
    with pytest.raises(ValueError):
        meetings.manager.start("x", language="fr")


HINDI_LINES = _segs((0.0, 4.0, "मेरा बाई आई आज सकूल नहीं गया"), (5.0, 9.0, "तुजे पता है क्यल मीटिं"),
                    (10.0, 14.0, "बजट कितना है"), (15.0, 19.0, "ठीक है"))


def _hindi_meeting(temp_audio: bool) -> tuple[int, list[int], Path]:
    from foresight import repass
    from foresight.config import settings

    mid = db.execute("INSERT INTO meetings(title, started_at, ended_at, status, language) VALUES (?,?,?,?,?)",
                     ("Hindi standup", time.time(), time.time() + 60, "done", "hi"))
    with db.tx() as conn:
        seg_ids = [conn.execute("INSERT INTO segments(meeting_id, t_start, t_end, speaker, text) VALUES (?,?,?,?,?)",
                                (mid, s["t_start"], s["t_end"], "Recording", s["text"])).lastrowid for s in HINDI_LINES]
    path = settings.recordings_dir / f"hindi-{mid}.wav"
    path.write_bytes(llm.wav_bytes(_tone(20.0)))
    media.adopt_import(mid, path, keep=True)
    repass.prepare(mid, temp_audio=temp_audio)
    indexer._upsert_source(f"meeting:{mid}:transcript", "Hindi standup — transcript", "transcript", meeting_id=mid, mime="md")
    return mid, seg_ids, path


def test_gemma_pass_rewrites_hindi_meetings(monkeypatch):
    from foresight import repass

    replies = iter(["Mera bhai aaj school nahi gaya",          # better: used
                    "",                                         # nothing heard: the live line stays
                    "I'm sorry, I can't transcribe this audio",  # a refusal: stays
                    "theek hai " * 40])                          # far longer than the line: a hallucination, stays
    heard = []
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: heard.append(messages[0]) or next(replies))
    reindexed = []
    monkeypatch.setattr(indexer, "add_meeting_text", lambda *a: reindexed.append(a))

    mid, seg_ids, path = _hindi_meeting(temp_audio=True)
    try:
        assert repass.run(mid) == {"status": "done", "done": 4, "total": 4, "changed": 1, "temp_audio": 1, "error": None}
        assert all(len(m["images"]) == 1 and "Latin letters" in m["content"] for m in heard)
        texts = [r["text"] for r in db.query("SELECT text FROM segments WHERE meeting_id = ? ORDER BY t_start", (mid,))]
        assert texts == ["Mera bhai aaj school nahi gaya"] + [s["text"] for s in HINDI_LINES[1:]]
        assert db.query("SELECT text, replaced_by FROM segment_versions WHERE segment_id = ?", (seg_ids[0],)) == [
            {"text": HINDI_LINES[0]["text"], "replaced_by": "gemma"}]
        assert reindexed == [(mid, "transcript", "Hindi standup — transcript")]
        # The audio was kept only for the pass.
        assert media.for_meeting(mid) == [] and not path.exists()
        assert repass.run(mid)["status"] == "done" and len(heard) == 4  # a finished pass is not repeated
    finally:
        meetings.delete(mid)


def test_gemma_pass_resumes_and_keeps_opted_in_audio(monkeypatch):
    from foresight import repass

    calls = []
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: calls.append(1) or "theek hai")
    monkeypatch.setattr(indexer, "add_meeting_text", lambda *a: None)
    mid, seg_ids, path = _hindi_meeting(temp_audio=False)
    try:
        db.execute("UPDATE transcript_passes SET status = 'running', done = 3 WHERE meeting_id = ?", (mid,))  # interrupted
        assert repass.run(mid)["changed"] == 1 and len(calls) == 1  # only the last line was left
        assert path.exists() and media.for_meeting(mid)[0]["url"]  # Keep meeting audio was on: the audio stays
    finally:
        meetings.delete(mid)


def test_import_a_hindi_recording(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from foresight import asr, repass
    from foresight.server import app

    class FakeEngine:
        def transcribe_file(self, path):
            yield 0.0, 2.0, "नमस्ते सब लोग"

        def transcribe(self, audio, prompt=None):  # an encrypted import is decoded in memory and cut by the VAD
            return "नमस्ते सब लोग"

    engines = []
    monkeypatch.setattr(asr.Transcriber, "for_language", classmethod(lambda cls, lang="en": engines.append(lang) or FakeEngine()))
    queued = []
    monkeypatch.setattr(repass, "enqueue", lambda mid: queued.append(mid))
    monkeypatch.setattr(indexer, "add_meeting_text", lambda *a: None)
    wav = llm.wav_bytes(_tone(2.0))
    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        bad = client.post("/api/meetings/import", files={"file": ("x.wav", wav, "audio/wav")}, data={"language": "fr"},
                          headers={"X-Foresight": "1"})
        assert bad.status_code == 400
        mid = client.post("/api/meetings/import", files={"file": ("namaste.wav", wav, "audio/wav")},
                          data={"title": "Namaste", "language": "hi"}, headers={"X-Foresight": "1"}).json()["id"]
        for _ in range(50):
            if client.get(f"/api/meetings/{mid}").json()["status"] == "done":
                break
            time.sleep(0.1)
        m = client.get(f"/api/meetings/{mid}").json()
    try:
        assert engines == ["hi"] and m["language"] == "hi" and m["segments"][0]["text"] == "नमस्ते सब लोग"
        assert m["transcript_pass"]["temp_audio"] == 1 and queued == [mid]
        assert m["media"] and m["media"][0]["url"]  # kept for Gemma's pass although Keep meeting audio is off
        from foresight import vault
        stored = db.query_one("SELECT path FROM media WHERE id = ?", (m["media"][0]["id"],))["path"]
        assert vault.is_encrypted(stored) and wav[44:4096] not in Path(stored).read_bytes()  # never plain on disk
    finally:
        meetings.delete(mid)


def test_stopping_can_leave_the_backlog_for_later():
    from foresight import asr

    capture = asr.LiveCapture(lambda *a: None, mic=True, system=False, language="hi")
    for t in (0.0, 15.0):
        capture.utterances.put(asr.Utterance("You", t, np.zeros(16000, dtype=np.float32)))
    assert capture.stop(transcribe_backlog=False) == 2 and capture.utterances.get_nowait() is None


def test_gemma_pass_fills_speech_live_transcription_missed(monkeypatch):
    from foresight import repass
    from foresight.config import settings

    # 30 s of audio with speech-like sound at 0–4 s, 10–14 s and 20–24 s; live transcription only got the first part.
    audio = np.zeros(30 * 16000, dtype=np.float32)
    for start in (0, 10, 20):
        audio[start * 16000:(start + 4) * 16000] = _tone(4.0)
    mid = db.execute("INSERT INTO meetings(title, started_at, ended_at, status, language) VALUES (?,?,?,?,?)",
                     ("Behind live", time.time(), time.time() + 30, "done", "hi"))
    first = db.execute("INSERT INTO segments(meeting_id, t_start, t_end, speaker, text) VALUES (?,?,?,?,?)",
                       (mid, 0.0, 4.0, "Recording", "पहली लाइन"))
    path = settings.recordings_dir / f"behind-{mid}.wav"
    path.write_bytes(llm.wav_bytes(audio))
    media.adopt_import(mid, path, keep=True)
    repass.prepare(mid, temp_audio=False)
    replies = iter(["pehli line", "doosri line jo live mein chhoot gayi", ""])  # the last gap was only noise
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: next(replies))
    monkeypatch.setattr(indexer, "add_meeting_text", lambda *a: None)
    try:
        result = repass.run(mid)
        lines = db.query("SELECT id, t_start, text FROM segments WHERE meeting_id = ? ORDER BY t_start", (mid,))
        assert [ln["text"] for ln in lines] == ["pehli line", "doosri line jo live mein chhoot gayi"]
        assert 9.0 < lines[1]["t_start"] < 10.5  # the gap line sits where the speech was
        assert result["status"] == "done" and result["changed"] == 2 and result["total"] == 3
        assert db.query("SELECT text FROM segment_versions WHERE segment_id = ?", (first,)) == [{"text": "पहली लाइन"}]
    finally:
        meetings.delete(mid)


def test_pass_length_limit_scales_with_the_audio():
    from foresight import repass

    assert repass.acceptable("", "a" * 500, seconds=15.0)        # a 15 s gap can hold a long line
    assert not repass.acceptable("", "a" * 700, seconds=15.0)    # ...but not this long
    assert not repass.acceptable("ok", "a" * 200)                 # short line, no duration: the old rule


def test_pass_drops_chatbot_replies_and_writes_latin_letters(monkeypatch):
    from foresight import repass

    assert not repass.acceptable("", "नमस्ते! मैं आपकी कैसे मदद कर सकता हूँ?", 4.0)
    assert not repass.acceptable("", "Hello! How can I help you today?", 4.0)
    assert repass.acceptable("", "aaj ka match bahut accha tha aur bowling bhi tagdi thi", 4.0)

    calls = []
    def fake_chat(messages, **kw):
        calls.append(messages)
        return "tujhe pata hai kal meeting kitne baje hai?"
    monkeypatch.setattr(llm, "chat", fake_chat)
    assert repass.in_latin_letters("तुझे पता है कल मीटिंग कितने बजे है?") == "tujhe pata hai kal meeting kitne baje hai?"
    assert "Latin letters" in calls[0][0]["content"]
    assert repass.in_latin_letters("already latin hinglish") == "already latin hinglish" and len(calls) == 1
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: "Here is a short summary.")  # not a transliteration
    assert repass.in_latin_letters("तुझे पता है कल मीटिंग कितने बजे है?") == "तुझे पता है कल मीटिंग कितने बजे है?"


def test_pass_skips_audio_without_speech(monkeypatch):
    """Clips without speech never reach Gemma: live lines written over music are removed, and gaps are not added."""
    from foresight.config import settings

    audio = np.zeros(30 * 16000, dtype=np.float32)
    for start in (0, 10, 20):
        audio[start * 16000:(start + 4) * 16000] = _tone(4.0)
    mid = db.execute("INSERT INTO meetings(title, started_at, ended_at, status, language) VALUES (?,?,?,?,?)",
                     ("Music", time.time(), time.time() + 30, "done", "hi"))
    with db.tx() as conn:
        conn.execute("INSERT INTO segments(meeting_id, t_start, t_end, speaker, text) VALUES (?,?,?,?,?)",
                     (mid, 0.0, 4.0, "Recording", "जाग़़ बाद्ट्ट्ट्ट्ट्ट्"))  # Whisper's text over a jingle
        conn.execute("INSERT INTO segments(meeting_id, t_start, t_end, speaker, text) VALUES (?,?,?,?,?)",
                     (mid, 20.0, 24.0, "Recording", "asli baat"))
    path = settings.recordings_dir / f"music-{mid}.wav"
    path.write_bytes(llm.wav_bytes(audio))
    media.adopt_import(mid, path, keep=True)
    repass.prepare(mid, temp_audio=False)
    # Only the part around 20 s is speech; 0–4 s and 10–14 s are music. The detector is asked, in order, about the
    # uncovered sound at 10 s (fill_gaps), then the lines at 0 s and 20 s (the pass).
    answers = iter([0.0, 0.0, 4.0])
    monkeypatch.setattr(repass, "speech_seconds", lambda pcm: next(answers))
    gemma = []
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: gemma.append(1) or "asli baat hai yeh")
    monkeypatch.setattr(indexer, "add_meeting_text", lambda *a: None)
    try:
        result = repass.run(mid)
        texts = [r["text"] for r in db.query("SELECT text FROM segments WHERE meeting_id = ? ORDER BY t_start", (mid,))]
        assert texts == ["asli baat hai yeh"] and len(gemma) == 1
        assert result["status"] == "done" and result["changed"] == 1
    finally:
        meetings.delete(mid)


# ---------------------------------------------------------------- reranker

def _hit(n: int, text: str) -> retrieval.Hit:
    return retrieval.Hit(n, n, f"doc{n}", f"uri{n}", "file", text, 1.0 / n, 0.5)


def test_rerank_puts_the_answering_passage_first(monkeypatch):
    hits = [_hit(1, "weather"), _hit(2, "the budget is 52000"), _hit(3, "budget history"), _hit(4, "vendors")]
    asked = []

    def fake_json(messages, **kw):
        asked.append(messages[1]["content"])
        return {"scores": [{"id": 1, "score": 0}, {"id": 2, "score": 3}, {"id": 3, "score": 2}, {"id": 4, "score": 1}]}

    monkeypatch.setattr(llm, "chat_json", fake_json)
    assert [h.title for h in retrieval.rerank("what is the budget?", hits, 3)] == ["doc2", "doc3", "doc4"]
    assert "Question: what is the budget?" in asked[0] and "[4] doc4" in asked[0]
    # Unrelated passages never pad the answer, however few relevant ones remain...
    monkeypatch.setattr(llm, "chat_json", lambda m, **kw: {"scores": [{"id": 1, "score": 0}, {"id": 2, "score": 3},
                                                                      {"id": 3, "score": 0}, {"id": 4, "score": 0}]})
    assert [h.title for h in retrieval.rerank("q", hits, 3)] == ["doc2"]
    # ...passages Gemma leaves out of its scores count as unrelated...
    monkeypatch.setattr(llm, "chat_json", lambda m, **kw: {"scores": [{"id": 3, "score": 2}]})
    assert [h.title for h in retrieval.rerank("q", hits, 3)] == ["doc3"]
    # ...and an off-topic question ("what is life") is left with no sources at all.
    monkeypatch.setattr(llm, "chat_json", lambda m, **kw: {"scores": [{"id": n, "score": 0} for n in range(1, 5)]})
    assert retrieval.rerank("what is life", hits, 3) == []
    # Any failure keeps the search order.
    monkeypatch.setattr(llm, "chat_json", lambda m, **kw: {})
    assert [h.title for h in retrieval.rerank("q", hits, 2)] == ["doc1", "doc2"]

    def broken(m, **kw):
        raise llm.LLMError("ollama down")
    monkeypatch.setattr(llm, "chat_json", broken)
    assert [h.title for h in retrieval.rerank("q", hits, 2)] == ["doc1", "doc2"]


def test_chat_reranks_when_enabled(monkeypatch):
    from foresight import rag
    from foresight.config import settings

    monkeypatch.setitem(runtime._load(), "rerank", True)
    seen = []
    monkeypatch.setattr(retrieval, "search", lambda q, k=6, **kw: [_hit(n, f"passage {n}") for n in range(1, k + 1)])
    monkeypatch.setattr(retrieval, "rerank", lambda q, hits, k: seen.append((q, len(hits), k)) or hits[::-1][:k])
    monkeypatch.setattr(llm, "chat_stream", lambda messages, **kw: iter(["ok [1]"]))
    start = next(e for e in rag.answer_stream("which passage?") if e["type"] == "start")
    assert seen == [("which passage?", settings.retrieve_k * 2, settings.retrieve_k)]
    assert start["citations"][0]["title"] == f"doc{settings.retrieve_k * 2}"  # reranked order reached the answer


def test_off_topic_question_gets_no_sources(monkeypatch):
    """"what is life" over meeting notes and READMEs: search always finds the least unrelated passages, and Gemma
    used to be handed them all and summarize them. Now it gets none, and is told so."""
    from foresight import rag

    monkeypatch.setitem(runtime._load(), "rerank", True)
    monkeypatch.setattr(retrieval, "search", lambda q, k=6, **kw: [_hit(n, f"Google Edge Foresight notes {n}")
                                                                   for n in range(1, k + 1)])
    monkeypatch.setattr(llm, "chat_json", lambda m, **kw: {"scores": [{"id": n, "score": 0} for n in range(1, 13)]})
    sent = []
    monkeypatch.setattr(llm, "chat_stream", lambda messages, **kw: sent.append(messages) or iter(["Not in your sources."]))
    events = list(rag.answer_stream("what is life"))
    start = next(e for e in events if e["type"] == "start")
    assert start["citations"] == []
    question = sent[0][-1]["content"]
    assert "(no matching sources were found)" in question and "Google Edge Foresight" not in question


# ---------------------------------------------------------------- date questions

def test_date_windows_in_questions():
    from foresight import dates

    now = datetime(2026, 10, 18, 15, 0)  # a Sunday afternoon

    def day(question):
        w = dates.parse_window(question, now)
        return None if w is None else (datetime.fromtimestamp(w[0]).strftime("%a %d"), round((w[1] - w[0]) / 86400, 2))

    assert day("What did we decide in last Tuesday's meeting?") == ("Tue 13", 1)
    assert day("pichle mangalvar ki meeting mein kya hua") == ("Tue 13", 1)
    assert day("summarise Friday's standup") == ("Fri 16", 1)
    assert day("what happened yesterday?") == ("Sat 17", 1)
    assert day("kal ki meeting ka summary do") == ("Sat 17", 1)
    assert day("parso kya decide hua") == ("Fri 16", 1)
    assert day("anything from 3 days ago?") == ("Thu 15", 1)
    assert day("budget talk on 12 October") == ("Mon 12", 1)
    assert day("budget talk on Oct 9th") == ("Fri 09", 1)
    assert day("the 2026-10-05 review") == ("Mon 05", 1)
    assert day("action items from last week") == ("Mon 05", 7)
    assert day("is hafte kitni meetings hui") == ("Mon 12", 7)
    assert day("what did Priya say this morning") == ("Sun 18", 0.5)
    assert day("on Sunday") == ("Sun 18", 1)                      # today is Sunday
    assert day("last Sunday's retro") == ("Sun 11", 1)            # ...and last Sunday is a week ago
    assert day("Who sat in the design review?") is None           # "sat" is not Saturday
    assert day("What is the Atlas marketing budget?") is None
    assert day("may I see the vendor list") is None


def test_chat_searches_the_meeting_from_that_day(monkeypatch):
    from foresight import dates, rag

    now = datetime(2026, 10, 18, 15, 0)
    monkeypatch.setattr(dates, "_now", lambda: now)
    tuesday, friday = datetime(2026, 10, 13, 10, 0).timestamp(), datetime(2026, 10, 16, 10, 0).timestamp()
    ids = []
    for title, started, text in (("Tuesday sync", tuesday, "We decided the launch budget is 52000 dollars."),
                                 ("Friday sync", friday, "We decided the launch budget stays at 40000 dollars.")):
        mid = db.execute("INSERT INTO meetings(title, started_at, ended_at, status) VALUES (?,?,?,?)",
                         (title, started, started + 1800, "done"))
        db.execute("INSERT INTO segments(meeting_id, t_start, t_end, speaker, text) VALUES (?,?,?,?,?)",
                   (mid, 0.0, 5.0, "Others", text))
        indexer._process(indexer._upsert_source(f"meeting:{mid}:transcript", f"{title} — transcript", "transcript",
                                                meeting_id=mid, mime="md"))
        ids.append(mid)
    sent = []
    monkeypatch.setattr(llm, "chat_stream", lambda messages, **kw: sent.append(messages) or iter(["52000 [1]"]))
    try:
        start = list(rag.answer_stream("What budget did we decide in last Tuesday's meeting?"))[0]  # read it all
        assert start["window"] == {"label": "Tuesday 13 October", "meetings": 1}
        assert start["citations"] and {c["title"] for c in start["citations"]} == {"Tuesday sync — transcript"}
        assert "Meetings recorded then:" in sent[-1][-1]["content"] and "Tuesday sync" in sent[-1][-1]["content"]
        # No meeting that day: say so, and search everything.
        start = list(rag.answer_stream("What budget did we decide on Wednesday?"))[0]
        assert start["window"]["meetings"] == 0 and len({c["title"] for c in start["citations"]}) >= 2
        assert "(no meetings were recorded then)" in sent[-1][-1]["content"]
    finally:
        for mid in ids:
            meetings.delete(mid)


def _index_folder(path) -> int:
    """Index a folder here and now. (A server started by an earlier test runs the real background worker, which
    would otherwise race this test for the queue.)"""
    queued = []
    original = indexer.enqueue
    indexer.enqueue = queued.append
    try:
        indexer.add_folder(str(path))
    finally:
        indexer.enqueue = original
    for sid in queued:
        indexer._process(sid)
    return len(queued)


def test_graph_links_similar_sources_and_their_labels(tmp_path):
    from fastapi.testclient import TestClient

    from foresight import graph
    from foresight.server import app

    docs = {"orion-pricing.md": "Orion pricing tiers: starter plan, growth plan and enterprise plan with a volume discount.",
            "orion-offer.md": "The Orion enterprise plan gets a volume discount; starter and growth plan pricing tiers stay.",
            "garden.txt": "Tomato seedlings need watering twice weekly. Compost the raised garden beds in spring."}
    for name, text in docs.items():
        (tmp_path / name).write_text(text, encoding="utf-8")
    _index_folder(tmp_path)
    ids = {r["title"]: r["id"] for r in db.query("SELECT id, title FROM sources WHERE title IN (?,?,?)", tuple(docs))}
    with db.tx() as conn:
        conn.execute("INSERT OR IGNORE INTO labels(name) VALUES ('Orion-map')")
        for sid in ids.values():
            conn.execute("INSERT INTO source_labels VALUES (?, (SELECT id FROM labels WHERE name = 'Orion-map'))", (sid,))

    out = graph.build(label="Orion-map")
    nodes = {n["id"]: n for n in out["nodes"]}
    assert set(nodes) == {*ids.values(), "label:Orion-map"}
    similar = {frozenset((e["a"], e["b"])) for e in out["edges"] if e["kind"] == "similar"}
    assert similar == {frozenset((ids["orion-pricing.md"], ids["orion-offer.md"]))}  # the garden note stays apart
    assert sorted(e["b"] for e in out["edges"] if e["kind"] == "label") == sorted(ids.values())
    assert nodes[ids["garden.txt"]]["labels"] == ["Orion-map"] and nodes["label:Orion-map"]["kind"] == "label"
    assert graph.build(label="No-such-label")["nodes"] == []

    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        assert client.get("/api/graph", params={"label": "Orion-map"}).json() == out
        everything = client.get("/api/graph").json()
        assert set(ids.values()) <= {n["id"] for n in everything["nodes"]}


def test_pictures_are_described_found_shown_and_seen(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from PIL import Image

    from foresight import images, rag
    from foresight.server import app

    looked_at = []

    def fake_chat(messages, **kw):  # Gemma describing the picture
        looked_at.append(messages[-1].get("images"))
        return "A whiteboard photo of the Zephyr roadmap. Text on it: Zephyr beta ships in March; owner Kofi."

    monkeypatch.setattr(llm, "chat", fake_chat)
    # The picture's own vector: what it looks like, which its description does not say.
    monkeypatch.setattr(llm, "embed_images", lambda pictures: fake_embed(["crimson poster artwork"] * len(pictures)))

    folder = tmp_path / "pictures"
    folder.mkdir()
    noise = np.random.default_rng(1).integers(0, 255, (200, 300, 3), dtype=np.uint8)
    Image.fromarray(noise).save(folder / "whiteboard.png")       # big enough to be knowledge
    Image.new("RGB", (16, 16), (200, 0, 0)).save(folder / "icon.png")  # an icon: skipped
    assert _index_folder(folder) == 1

    src = db.query_one("SELECT * FROM sources WHERE title = 'whiteboard.png'")
    assert src["status"] == "ready" and len(looked_at) == 1 and len(looked_at[0]) == 1
    assert "Zephyr beta ships in March" in db.query_one("SELECT text FROM chunks WHERE source_id = ?", (src["id"],))["text"]
    pic = db.query_one("SELECT * FROM media WHERE source_id = ?", (src["id"],))
    assert (pic["kind"], pic["owned"], pic["mime"], pic["status"]) == ("image", 0, "image/png", "ready")
    assert images.thumb_path(pic["id"]).is_file()
    assert db.query_one("SELECT COUNT(*) AS n FROM media_items WHERE media_id = ? AND t_start IS NULL", (pic["id"],))["n"] == 1

    by_words = retrieval.search("when does the Zephyr beta ship", k=3)
    assert by_words[0].title == "whiteboard.png" and by_words[0].media["kind"] == "image"
    assert by_words[0].media["thumb"] == f"/api/media/{pic['id']}/thumb"
    by_look = retrieval.search("crimson poster artwork", k=3)  # only the picture's vector knows this
    assert src["id"] in [h.source_id for h in by_look[:2]] and all(h.media["kind"] == "image" for h in by_look
                                                                   if h.source_id == src["id"])
    assert src["id"] not in [h.source_id for h in retrieval.search("vendor shortlist owner", k=3)
                             if h.kind == "image"]  # an unrelated question does not bring the picture along
    words_only = {h.source_id: h.similarity for h in retrieval.search("crimson poster artwork", k=50, with_media=False)}
    assert words_only.get(src["id"], 0.0) < 0.5  # its description does not match; its look does

    indexer._process(src["id"])  # unchanged picture: Gemma does not look again
    assert len(looked_at) == 1

    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        full = client.get(f"/api/media/{pic['id']}")
        assert full.status_code == 200 and full.headers["content-type"] == "image/png"
        assert client.get(f"/api/media/{pic['id']}/thumb").headers["content-type"] == "image/jpeg"
        assert client.get(f"/api/sources/{src['id']}/text").json()["picture"]["media_id"] == pic["id"]
        found = client.get("/api/search", params={"q": "Zephyr roadmap whiteboard", "modality": "image"}).json()
        assert found and all(h["media"]["kind"] == "image" for h in found)
        assert not any((h.get("media") or {}).get("kind") == "image" for h in
                       client.get("/api/search", params={"q": "Zephyr roadmap whiteboard", "modality": "audio"}).json())

    sent = []
    monkeypatch.setattr(llm, "chat_stream", lambda messages, **kw: sent.append(messages) or iter(["March [1]"]))
    start = list(rag.answer_stream("When does the Zephyr beta ship?"))[0]
    n = next(c["n"] for c in start["citations"] if c["source_id"] == src["id"])
    assert start["looked"] == [n] and start["listened"] == []
    user = sent[-1][-1]
    assert len(user["images"]) == 1 and "Pictures are attached" in user["content"]

    indexer.remove_source(src["id"])
    assert db.query_one("SELECT id FROM media WHERE id = ?", (pic["id"],)) is None
    assert not images.thumb_path(pic["id"]).exists()


def test_pictures_wait_for_meetings_and_can_be_left_out(tmp_path, monkeypatch):
    from PIL import Image

    described = []
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: described.append(1) or "A photo of a whiteboard.")
    monkeypatch.setattr(llm, "embed_images", lambda pictures: fake_embed(["whiteboard"] * len(pictures)))
    noise = np.random.default_rng(2).integers(0, 255, (160, 240, 3), dtype=np.uint8)
    Image.fromarray(noise).save(tmp_path / "board.jpg")
    (tmp_path / "agenda.txt").write_text("Agenda: review the Kestrel launch checklist.", encoding="utf-8")

    monkeypatch.setitem(runtime._load(), "index_pictures", False)
    assert _index_folder(tmp_path) == 1  # the text file only
    assert db.query_one("SELECT id FROM sources WHERE title = 'board.jpg'") is None

    monkeypatch.setitem(runtime._load(), "index_pictures", True)
    queued = []
    monkeypatch.setattr(indexer, "enqueue", queued.append)
    indexer.scan_folder(db.query_one("SELECT folder_id FROM sources WHERE title = 'agenda.txt'")["folder_id"])
    board = db.query_one("SELECT id FROM sources WHERE title = 'board.jpg'")["id"]
    assert queued == [board]

    recording = [True]
    monkeypatch.setattr(indexer, "_recording", lambda: recording[0])
    monkeypatch.setattr(indexer, "MEETING_POLL_SECONDS", 0.01)
    monkeypatch.setattr(indexer, "_wait_for_meeting_end", lambda: None)  # the test ends the meeting itself
    indexer._handle(board)  # a meeting is recording: the picture waits, Gemma is not called
    assert described == [] and board in indexer._deferred
    recording[0] = False
    queued.clear()
    indexer._after_meeting()  # the meeting ended: back in the queue
    assert queued == [board] and not indexer._deferred
    indexer._handle(board)
    assert described == [1]
    assert db.query_one("SELECT status FROM sources WHERE id = ?", (board,))["status"] == "ready"


def test_voice_memos_become_searchable_and_playable(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from foresight import asr, listen, memos, rag
    from foresight.server import app

    heard = []

    class Engine:  # what the speech engine heard in the memo
        def transcribe_file(self, path):
            heard.append(Path(path).name)
            yield 0.4, 4.0, "Remember to book the Lisbon venue for the offsite."
            yield 5.0, 9.5, "And ask Tomas about the catering budget."

    monkeypatch.setattr(asr.Transcriber, "for_language", classmethod(lambda cls, language="en": Engine()))
    # Audio moments: the sound of each window (here: what was said in it).
    windows = []
    monkeypatch.setattr(llm, "embed_audio", lambda wavs: windows.append(len(wavs)) or fake_embed(
        ["Lisbon venue offsite catering budget Tomas"] * len(wavs)))

    folder = tmp_path / "memos"
    folder.mkdir()
    memo = folder / "offsite idea.wav"
    memo.write_bytes(llm.wav_bytes(np.concatenate([_tone(5.0), _tone(5.0, 330.0)])))
    (folder / "ding.wav").write_bytes(llm.wav_bytes(_tone(0.3)))  # a notification sound: skipped
    assert _index_folder(folder) == 1

    src = db.query_one("SELECT * FROM sources WHERE title = 'offsite idea.wav'")
    assert src["status"] == "ready" and heard == ["offsite idea.wav"]
    text = db.query_one("SELECT text FROM chunks WHERE source_id = ?", (src["id"],))["text"]
    assert text.startswith("Voice memo: offsite idea.wav (00:10)") and "[00:05] And ask Tomas about the catering" in text
    rec = db.query_one("SELECT * FROM media WHERE source_id = ?", (src["id"],))
    assert (rec["kind"], rec["owned"], rec["mime"], rec["status"]) == ("audio", 0, "audio/wav", "ready")
    assert abs(rec["duration_s"] - 10.0) < 0.1 and sum(windows) >= 1
    moments_rows = db.query("SELECT t_start, t_end, text FROM media_items WHERE media_id = ?", (rec["id"],))
    assert moments_rows and "Lisbon" in moments_rows[0]["text"]

    hits = retrieval.search("who do I ask about the catering budget for the Lisbon offsite", k=3)
    top = next(h for h in hits if h.source_id == src["id"])
    assert top.media and top.media["url"] == f"/api/media/{rec['id']}" and top.media["t_start"] < 9.5

    picked = listen.moment(rec["id"], 4.0, 10.0)
    assert picked["title"].startswith("offsite idea.wav — audio") and "Tomas" in picked["text"]
    assert memos.lines_between(src["id"], 0, 3) == ["Remember to book the Lisbon venue for the offsite."]

    indexer._process(src["id"])  # unchanged file: not transcribed again
    assert heard == ["offsite idea.wav"]

    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        played = client.get(f"/api/media/{rec['id']}", headers={"Range": "bytes=0-99"})
        assert played.status_code == 206 and played.headers["content-type"] == "audio/wav"
        assert client.get(f"/api/sources/{src['id']}/text").json()["audio"]["media_id"] == rec["id"]

    sent = []
    monkeypatch.setattr(llm, "chat_stream", lambda messages, **kw: sent.append(messages) or iter(["Tomas [1]"]))
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: "And ask Tomas about the catering budget.")  # its hearing
    start = list(rag.answer_stream("Who should I ask about the Lisbon offsite catering?", listen_audio=True))[0]
    n = next(c["n"] for c in start["citations"] if c["source_id"] == src["id"])
    assert n in start["listened"] and "Heard in the recording, word for word" in sent[-1][-1]["content"]  # Gemma heard it

    indexer.remove_source(src["id"])
    assert db.query_one("SELECT id FROM media WHERE id = ?", (rec["id"],)) is None
    assert memo.is_file()  # the user's own file is never touched


def test_files_at_rest_are_encrypted_and_read_back(tmp_path):
    from foresight import vault

    assert vault.available() and vault.key() not in vault.KEY_FILE.read_bytes()  # wrapped by DPAPI, not stored plain
    data = np.random.default_rng(3).integers(0, 256, vault.CHUNK * 2 + 1234, dtype=np.uint8).tobytes()
    path = tmp_path / "blob.bin"
    vault.write_bytes(path, data)
    raw = path.read_bytes()
    assert vault.is_encrypted(path) and data[:64] not in raw and vault.plain_size(path) == len(data)
    assert vault.read_bytes(path) == data
    with vault.Reader(path) as r:  # random access across chunk boundaries
        r.seek(vault.CHUNK - 10)
        assert r.read(30) == data[vault.CHUNK - 10:vault.CHUNK + 20]
        r.seek(-5, 2)
        assert r.read() == data[-5:]

    tampered = bytearray(raw)
    tampered[vault.HEADER + 100] ^= 1
    (tmp_path / "tampered.bin").write_bytes(bytes(tampered))
    with pytest.raises(vault.VaultError):
        vault.read_bytes(tmp_path / "tampered.bin")
    stride = vault.CHUNK + vault.TAG
    (tmp_path / "cut.bin").write_bytes(raw[:vault.HEADER + stride + 7])  # cut inside a chunk: refused
    with pytest.raises(vault.VaultError):
        vault.read_bytes(tmp_path / "cut.bin")
    (tmp_path / "crashed.bin").write_bytes(raw[:vault.HEADER + stride])  # a recording cut short at a chunk boundary
    assert vault.read_bytes(tmp_path / "crashed.bin") == data[:vault.CHUNK]

    plain = tmp_path / "old.txt"
    plain.write_text("from before encryption", encoding="utf-8")
    assert vault.read_bytes(plain) == b"from before encryption"  # plain files still read
    assert vault.encrypt_existing([tmp_path]) == 1 and vault.is_encrypted(plain)
    assert vault.read_bytes(plain) == b"from before encryption" and vault.encrypt_file(plain) is False


def test_uploads_are_stored_encrypted_and_served_decrypted(monkeypatch):
    import io

    import docx
    from fastapi.testclient import TestClient

    from foresight import vault
    from foresight.server import app

    d = docx.Document()
    d.add_paragraph("The Heron project budget is 77,000 dollars.")
    buf = io.BytesIO()
    d.save(buf)
    queued = []
    monkeypatch.setattr(indexer, "enqueue", queued.append)
    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        up = client.post("/api/sources/upload", files=[("files", ("heron.docx", buf.getvalue()))],
                         headers={"X-Foresight": "1"}).json()
        sid = up["added"][0]
        stored = db.query_one("SELECT uri FROM sources WHERE id = ?", (sid,))["uri"]
        assert vault.is_encrypted(stored)
        indexer._process(sid)  # parsed from the decrypted bytes, in memory
        assert "77,000" in client.get(f"/api/sources/{sid}/text").json()["chunks"][0]

        # Kept meeting audio: the player's byte ranges come back decrypted.
        mid = _new_meeting("Range test")
        store = media.AudioStore(mid, ["You"])
        store.feed("You", _tone(2.0))
        store.close()
        rec = db.query_one("SELECT id, path FROM media WHERE meeting_id = ?", (mid,))
        assert vault.is_encrypted(rec["path"])
        whole = vault.read_bytes(rec["path"])
        part = client.get(f"/api/media/{rec['id']}", headers={"Range": "bytes=100-299"})
        assert part.status_code == 206 and part.content == whole[100:300]
        assert part.headers["content-range"] == f"bytes 100-299/{len(whole)}"
        assert client.get(f"/api/media/{rec['id']}", headers={"Range": "bytes=-50"}).content == whole[-50:]
        assert client.get(f"/api/media/{rec['id']}").content == whole
        assert client.get(f"/api/media/{rec['id']}", headers={"Range": f"bytes={len(whole)}-"}).status_code == 416
        assert media.decode_slice(rec["path"], 0.5, 1.5).shape[0] == pytest.approx(media.RATE, abs=400)
    meetings.delete(mid)
    indexer.remove_source(sid)


def test_first_run_check_says_what_is_missing_and_how_to_fix_it(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from foresight import asr, doctor
    from foresight.config import settings
    from foresight.server import app

    monkeypatch.setattr(llm, "status", lambda: {"ok": True, "version": "0.40", "gen_ready": False, "embed_ready": True})
    monkeypatch.setattr(asr, "list_devices", lambda: {"error": "no microphone"})
    hub = tmp_path / "hub"
    monkeypatch.setattr("huggingface_hub.constants.HF_HUB_CACHE", str(hub))
    whisper = hub / f"models--{doctor.whisper_repo().replace('/', '--')}" / "snapshots" / "abc"
    whisper.mkdir(parents=True)
    (whisper / "model.bin").write_bytes(b"weights")  # Whisper is there; Parakeet is not

    r = doctor.run()
    by_id = {c["id"]: c for c in r["checks"]}
    assert not r["ok"] and set(r["missing"]) == {"gen_model", "speech_en"}
    assert by_id["gen_model"]["fix"] == f"ollama pull {settings.gen_model}" and by_id["embed_model"]["ok"]
    assert by_id["speech_hi"]["ok"] and "setup.bat" in by_id["speech_en"]["fix"]
    assert not by_id["audio"]["ok"] and not by_id["audio"]["required"]  # a missing mic does not block setup
    assert by_id["data"]["ok"] and by_id["vault"]["ok"]
    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        assert client.get("/api/setup").json()["missing"] == r["missing"]


def test_tray_icon_image():
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    import run

    img = run.tray_image(64)
    assert img.size == (64, 64) and img.mode == "RGBA"
    assert img.getpixel((0, 0))[3] == 0 and img.getpixel((32, 32))[3] == 255  # a round orb on a clear background


def test_gemma_answers_from_what_it_hears_not_a_wrong_transcript(tmp_path, monkeypatch):
    from foresight import asr, rag

    class Engine:  # speech recognition got the amount wrong
        def transcribe_file(self, path):
            yield 0.2, 6.0, "The deposit for the Osprey venue is €400, due on 12 November."

    monkeypatch.setattr(asr.Transcriber, "for_language", classmethod(lambda cls, language="en": Engine()))
    monkeypatch.setattr(llm, "embed_audio", lambda wavs: fake_embed(["Osprey venue deposit"] * len(wavs)))
    folder = tmp_path / "osprey"
    folder.mkdir()
    (folder / "osprey deposit.wav").write_bytes(llm.wav_bytes(_tone(6.0)))
    _index_folder(folder)
    src = db.query_one("SELECT id FROM sources WHERE title = 'osprey deposit.wav'")["id"]

    heard = ["The deposit for the Osprey venue is four thousand euros, due on the twelfth of November."]
    monkeypatch.setattr(llm, "chat", lambda messages, **kw: heard[0])
    sent = []
    monkeypatch.setattr(llm, "chat_stream", lambda messages, **kw: sent.append(messages) or iter(["Four thousand euros [1]."]))
    try:
        events = list(rag.answer_stream("How much is the deposit for the Osprey venue?", listen_audio=True))
        n = next(c["n"] for c in events[0]["citations"] if c["source_id"] == src)
        assert n in events[0]["listened"]
        assert {"type": "heard", "texts": {str(n): heard[0]}} in events
        content = sent[-1][-1]["content"]
        # What was heard replaces the transcript lines the clip covers; the rest of the block stays.
        assert f"[{n}] osprey deposit.wav\nVoice memo: osprey deposit.wav (00:06)\n\n[00:00–00:06] Heard in the " \
               f"recording, word for word: {heard[0]}" in content
        assert "€400" not in content and "Answer from what was heard" in content
        assert not any(base64_wav(x).startswith(b"RIFF") for x in sent[-1][-1].get("images", []))  # heard once already

        for reply in ("NO_SPEECH", "Namaste! Main aapki kaise madad kar sakta hoon?"):  # silence, a chatbot reply
            heard[0] = reply
            events = list(rag.answer_stream("How much is the deposit for the Osprey venue?", listen_audio=True))
            assert not any(e["type"] == "heard" for e in events)
            assert "Heard in the recording" not in sent[-1][-1]["content"] and "trust the audio" in sent[-1][-1]["content"]
            assert any(base64_wav(x).startswith(b"RIFF") for x in sent[-1][-1]["images"])  # not heard: clips go along

        def offline(messages, **kw):
            raise llm.LLMError("Ollama is not running")
        monkeypatch.setattr(llm, "chat", offline)  # a failed hearing still answers, from the transcript
        assert list(rag.answer_stream("How much is the deposit for the Osprey venue?", listen_audio=True))[-1]["type"] == "done"
    finally:
        indexer.remove_source(src)


def test_heard_words_replace_the_transcript_lines_their_clip_covers():
    from foresight import rag

    chunk = ("Meeting: Atlas sync (2026-10-18 13:08)\n\n[00:00] Recording: Okay, let's start.\n"
             "[00:12] Recording: The budget is $5,200.\nIt was approved on Tuesday.\n[00:31] Recording: Next item.")
    out = rag.with_heard(chunk, (12.0, 30.0), "The budget is fifty two thousand dollars. It was approved on Tuesday.")
    assert out == ("Meeting: Atlas sync (2026-10-18 13:08)\n\n[00:00] Recording: Okay, let's start.\n"
                   "[00:12–00:30] Heard in the recording, word for word: The budget is fifty two thousand dollars. "
                   "It was approved on Tuesday.\n[00:31] Recording: Next item.")
    assert rag.with_heard("no timestamps here", (3.0, 9.0), "what was said") == \
        "[00:03–00:09] Heard in the recording, word for word: what was said"


def test_set_up_now_installs_ollama_then_pulls_the_models(monkeypatch):
    from fastapi.testclient import TestClient

    from foresight import doctor, provision
    from foresight.config import settings
    from foresight.server import app

    state = {"ok": False, "gen_ready": False, "embed_ready": False}
    monkeypatch.setattr(llm, "status", lambda: dict(state))
    monkeypatch.setattr(llm, "check_media", lambda: {})
    monkeypatch.setattr(provision, "ollama_exe", lambda: None)
    assert provision.plan() == ["install_ollama", f"pull:{settings.gen_model}", f"pull:{settings.embed_model}"]
    monkeypatch.setattr(doctor, "INSTALLED", True)  # as ForesightWing.exe: the fixes point at the button
    r = doctor.run(devices=False)
    assert r["download_gb"] == 6.6 and {c["fix"] for c in r["checks"] if c["id"] in ("ollama", "gen_model")} == {"Press Set up now."}

    done = []
    monkeypatch.setattr(provision, "install_ollama", lambda: done.append("install"))
    monkeypatch.setattr(provision, "start_ollama", lambda: done.append("start") or state.update(ok=True) or True)
    monkeypatch.setattr(provision, "pull", lambda model: done.append(model))
    with TestClient(app, base_url="http://127.0.0.1:8765") as client:
        assert client.post("/api/setup/run").status_code == 403  # a button press from this app only
        client.post("/api/setup/run", headers={"X-Foresight": "1"})
        for _ in range(50):
            if not client.get("/api/setup/progress").json()["running"]:
                break
            time.sleep(0.05)
        job = client.get("/api/setup/progress").json()
    assert done == ["install", "start", settings.gen_model, settings.embed_model]
    assert job["finished"] and job["error"] is None and job["step"] == "Ready"


def test_an_unsigned_ollama_installer_is_never_run(tmp_path, monkeypatch):
    import subprocess

    from foresight import provision

    fake = tmp_path / "not-ollama.exe"
    fake.write_bytes(b"MZ" + b"\0" * 64)
    assert provision.signed_by_ollama(fake) is False  # Windows' own signature check

    def download(url, dest, step):
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(b"MZ tampered")

    ran = []
    monkeypatch.setattr(provision, "_download", download)
    monkeypatch.setattr(subprocess, "run", lambda *a, **kw: ran.append(a))
    monkeypatch.setattr(provision, "signed_by_ollama", lambda path: False)
    with pytest.raises(RuntimeError, match="not signed by Ollama"):
        provision.install_ollama()
    assert ran == [] and not (provision.settings.data_dir / "downloads" / "OllamaSetup.exe").exists()


def test_installer_replaces_only_its_own_folder(tmp_path):
    from foresight import install

    payload = tmp_path / "payload"
    payload.mkdir()
    (payload / install.EXE).write_bytes(b"new")
    target = tmp_path / "Programs" / "ForesightWing"
    install.place(payload, target)
    assert (target / install.EXE).read_bytes() == b"new" and not payload.exists()

    update = tmp_path / "update"
    update.mkdir()
    (update / install.EXE).write_bytes(b"newer")
    install.place(update, target)  # an earlier install is replaced
    assert (target / install.EXE).read_bytes() == b"newer"

    someone_elses = tmp_path / "Documents"
    someone_elses.mkdir()
    (someone_elses / "thesis.docx").write_bytes(b"precious")
    again = tmp_path / "again"
    again.mkdir()
    (again / install.EXE).write_bytes(b"x")
    with pytest.raises(RuntimeError, match="holds other files"):
        install.place(again, someone_elses)
    assert (someone_elses / "thesis.docx").read_bytes() == b"precious"

    link = tmp_path / "Start Menu" / "Foresight Wing.lnk"
    install.make_shortcut(link, target / install.EXE)
    assert link.is_file()


@pytest.mark.skipif(sys.platform != "win32", reason="Windows file locking")
def test_uninstall_waits_for_the_app_to_exit(tmp_path):
    """The clean-up used to run 2 s after it started, while the uninstaller's "was removed" box was still open: the
    files the exe had loaded were locked, so the exe and its DLLs were left behind, and that exe no longer started."""
    import subprocess

    from foresight import install

    target = tmp_path / "Programs" / "ForesightWing"
    target.mkdir(parents=True)
    (target / install.EXE).write_bytes(b"x")
    dll = target / "python313.dll"
    dll.write_bytes(b"x")
    app = subprocess.Popen([sys.executable, "-c", f"import time; f = open({str(dll)!r}, 'rb'); time.sleep(60)"])
    try:
        cleanup = subprocess.Popen(install._remove_after_exit(target, app.pid))
        time.sleep(4)  # past the old 2 s: nothing may go while the app still runs
        assert (target / install.EXE).exists() and dll.exists()
    finally:
        app.kill()
    cleanup.wait(timeout=30)
    assert not target.exists()


def test_starts_without_a_console(tmp_path, monkeypatch):
    """ForesightWing.exe started from the Start menu has no console (sys.stdout and sys.stderr are None); uvicorn's log
    setup then crashed asking stdout.isatty(). Output now goes to foresight.log in the data folder."""
    import uvicorn

    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    import run

    monkeypatch.setattr(sys, "stdout", None)
    monkeypatch.setattr(sys, "stderr", None)
    run.console_or_log(tmp_path)
    try:
        uvicorn.Config(lambda scope, receive, send: None, log_level="warning")  # raised ValueError before
        print("hello from the tray")
        assert "hello from the tray" in (tmp_path / "foresight.log").read_text(encoding="utf-8")
    finally:
        sys.stdout.close()


def test_opens_in_its_own_window(tmp_path, monkeypatch):
    import subprocess
    import webbrowser

    from foresight import window

    started, browsed = [], []
    monkeypatch.setattr(subprocess, "Popen", lambda cmd, **kw: started.append(cmd))
    monkeypatch.setattr(webbrowser, "open", browsed.append)
    edge = tmp_path / "msedge.exe"
    monkeypatch.setattr(window, "edge", lambda: edge)
    assert window.open_window("http://127.0.0.1:8765/", tmp_path / "window") is True
    cmd = started[0]
    assert cmd[0] == str(edge) and "--app=http://127.0.0.1:8765/" in cmd  # an app window: no tabs, no address bar
    assert f"--user-data-dir={tmp_path / 'window'}" in cmd and not browsed  # its own profile, apart from the user's
    assert "--disable-extensions" in cmd and "--disable-sync" in cmd  # synced extensions could read transcripts and notes
    monkeypatch.setattr(window, "edge", lambda: None)  # no Edge: the default browser
    assert window.open_window("http://127.0.0.1:8765/", tmp_path / "window") is False and browsed == ["http://127.0.0.1:8765/"]

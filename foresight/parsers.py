"""Text extraction for every supported knowledge format.

pdf, docx, pptx, xlsx, odt, rtf, epub, html, md, txt (+ csv/json/xml as plain text).
Images and audio files are supported files too: Gemma describes images (images.py) and audio is transcribed
(memos.py) instead of parsed.
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup

from . import vault

TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".json", ".log", ".yaml", ".yml", ".xml", ".tsv"}
PARSED_EXTS = {".pdf", ".docx", ".pptx", ".xlsx", ".odt", ".rtf", ".epub", ".html", ".htm"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
AUDIO_EXTS = {".mp3", ".m4a", ".wav", ".ogg", ".opus", ".flac", ".aac"}
SUPPORTED_EXTS = TEXT_EXTS | PARSED_EXTS | IMAGE_EXTS | AUDIO_EXTS
MAX_FILE_BYTES = 80 * 1024 * 1024
MIN_FOLDER_IMAGE_BYTES = 12 * 1024  # smaller pictures in watched folders are icons and UI art, not knowledge
MIN_FOLDER_AUDIO_BYTES = 48 * 1024  # smaller audio files are notification and UI sounds


class ParseError(RuntimeError):
    pass


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_EXTS


def is_image(path: Path) -> bool:
    return path.suffix.lower() in IMAGE_EXTS


def is_audio(path: Path) -> bool:
    return path.suffix.lower() in AUDIO_EXTS


def _src(path: Path):
    """The file for a parser library: its path, or its decrypted bytes in memory when Foresight Wing encrypted it (uploads)."""
    return io.BytesIO(vault.read_bytes(path)) if vault.is_encrypted(path) else str(path)


def _read_text(path: Path) -> str:
    raw = vault.read_bytes(path)
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(_src(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Page {i}]\n{text}")
    return "\n\n".join(pages)


def _docx(path: Path) -> str:
    import docx

    document = docx.Document(_src(path))
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            parts.append("#" * min(int(level), 6) + " " + text)
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _xml_text(xml: bytes, tag_pattern: str) -> list[str]:
    return [m.group(1) for m in re.finditer(tag_pattern, xml.decode("utf-8", "replace"), re.S)]


def _pptx(path: Path) -> str:
    with zipfile.ZipFile(_src(path)) as z:
        slides = sorted(
            (n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)),
        )
        out = []
        for i, name in enumerate(slides, 1):
            texts = _xml_text(z.read(name), r"<a:t>(.*?)</a:t>")
            if texts:
                out.append(f"[Slide {i}]\n" + "\n".join(BeautifulSoup(t, "html.parser").get_text() for t in texts))
        return "\n\n".join(out)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(_src(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells).rstrip(" |"))
            if len(rows) >= 5000:
                break
        if rows:
            out.append(f"[Sheet {ws.title}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(out)


def _odt(path: Path) -> str:
    with zipfile.ZipFile(_src(path)) as z:
        soup = BeautifulSoup(z.read("content.xml"), "xml")
    return "\n".join(p.get_text(" ", strip=True) for p in soup.find_all(["text:h", "text:p"]) if p.get_text(strip=True))


def _rtf(path: Path) -> str:
    text = _read_text(path)
    text = re.sub(r"\\par[d]?", "\n", text)
    text = re.sub(r"\\'[0-9a-fA-F]{2}", "", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = re.sub(r"[{}]", "", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def html_to_text(html: str | bytes) -> tuple[str, str]:
    """Return (title, readable text) for an HTML document."""
    soup = BeautifulSoup(html, "lxml")
    title = (soup.title.string or "").strip() if soup.title and soup.title.string else ""
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "svg"]):
        tag.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    lines = []
    for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "td", "th", "blockquote"]):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            text = "#" * int(el.name[1]) + " " + text
        elif el.name == "li":
            text = "- " + text
        lines.append(text)
    body = "\n".join(lines) if lines else root.get_text("\n", strip=True)
    return title, body


def _epub(path: Path) -> str:
    with zipfile.ZipFile(_src(path)) as z:
        names = sorted(n for n in z.namelist() if n.lower().endswith((".xhtml", ".html", ".htm")))
        return "\n\n".join(html_to_text(z.read(n))[1] for n in names)


def extract(path: Path) -> str:
    """Extract plain text from a supported file. Raises ParseError for unreadable input."""
    ext = path.suffix.lower()
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ParseError(f"file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB")
    try:
        if ext in TEXT_EXTS:
            return _read_text(path)
        if ext == ".pdf":
            return _pdf(path)
        if ext == ".docx":
            return _docx(path)
        if ext == ".pptx":
            return _pptx(path)
        if ext == ".xlsx":
            return _xlsx(path)
        if ext == ".odt":
            return _odt(path)
        if ext == ".rtf":
            return _rtf(path)
        if ext == ".epub":
            return _epub(path)
        if ext in (".html", ".htm"):
            return html_to_text(vault.read_bytes(path))[1]
    except ParseError:
        raise
    except Exception as exc:  # third-party parsers raise many unrelated types
        raise ParseError(f"{type(exc).__name__}: {exc}") from exc
    raise ParseError(f"unsupported file type {ext}")
